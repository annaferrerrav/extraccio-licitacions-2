#!/usr/bin/env python3
"""
Consolida els fitxers files.jsonl + control.json de cada carpeta d'output/
en un únic Excel, una fila per línia extreta.

Pensat per créixer: es pot executar cada vegada que s'acaba un expedient
nou i tornarà a generar l'Excel amb tots els que hi hagi a output/ en
aquell moment.

Ús:
    python build_excel.py [directori_output] [fitxer_sortida.xlsx] [carpetes]

    carpetes:
        "all"                  -> totes les subcarpetes de output/ (per defecte)
        "4161038,9252209"      -> només aquests identificadors, separats per comes
        "4161038 9252209"      -> també accepta espais

Exemples:
    python build_excel.py
    python build_excel.py output extraccio_licitacions.xlsx all
    python build_excel.py output prova_dos.xlsx 4161038,9252209
    python build_excel.py output opus_repesca.xlsx 2222,3333,4444
"""

import json
import sys
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Ordre i etiquetes de les 13 columnes obligatòries, tal com al prompt.
CORE = [
    ("carpeta_analitzada", "Carpeta analitzada"),
    ("expedient", "Expedient"),
    ("proveidor", "Proveïdor"),
    ("descripcio", "Descripció"),
    ("referencia", "Referència"),
    ("model", "Model"),
    ("marca", "Marca"),
    ("preu_maxim_licitacio_unitari", "Preu màxim de licitació unitari"),
    ("quantitat", "Quantitat"),
    ("preu_unitari_sense_iva", "Preu unitari sense IVA"),
    ("import_total_sense_iva", "Import total sense IVA"),
    ("lot", "Lot"),
    ("estat", "Adjudicat / Oferta / Desestimat"),
]
NOTES_KEY = "notes"
NOTES_LABEL = "Notes"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=10)
ALT_FILL = PatternFill("solid", fgColor="F2F4F8")
CONF_COLOR = {"Alta": "1E7B34", "Mitjana": "B26B00", "Baixa": "B02418"}


def carrega_expedient(carpeta: pathlib.Path):
    """Retorna (files, confianca_per_linia) d'una carpeta output/<CARPETA>/."""
    jsonl = carpeta / "files.jsonl"
    control = carpeta / "control.json"
    if not jsonl.exists():
        return [], {}

    files = []
    for linia in jsonl.read_text(encoding="utf-8").splitlines():
        linia = linia.strip()
        if not linia:
            continue
        try:
            files.append(json.loads(linia))
        except json.JSONDecodeError:
            files.append({"_error_json": linia})

    confianca = {}
    if control.exists():
        try:
            d = json.loads(control.read_text(encoding="utf-8"))
            for x in d.get("confianca_per_fila", []):
                confianca[x.get("linia")] = x
        except json.JSONDecodeError:
            pass

    return files, confianca


def columnes_extra_ordenades(totes_les_files):
    """Claus que apareixen a les files, fora del CORE i de notes, en ordre d'aparició."""
    core_keys = {k for k, _ in CORE} | {NOTES_KEY}
    vistes = []
    for f in totes_les_files:
        for k in f.keys():
            if k not in core_keys and k not in vistes:
                vistes.append(k)
    return vistes


def main():
    base = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path("output")
    sortida = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else pathlib.Path("extraccio_licitacions.xlsx")
    seleccio = sys.argv[3] if len(sys.argv) > 3 else "all"

    if not base.exists():
        sys.exit(f"No existeix el directori {base}")

    totes = sorted([c for c in base.iterdir() if c.is_dir()])
    if not totes:
        sys.exit(f"No hi ha cap subcarpeta dins de {base}")

    if seleccio.strip().lower() == "all":
        carpetes = totes
    else:
        volgudes = [x.strip() for x in seleccio.replace(",", " ").split() if x.strip()]
        per_nom = {c.name: c for c in totes}
        carpetes = []
        no_trobades = []
        for nom in volgudes:
            if nom in per_nom:
                carpetes.append(per_nom[nom])
            else:
                no_trobades.append(nom)
        if no_trobades:
            print(f"Avís: no existeixen a {base}: {', '.join(no_trobades)}")
        if not carpetes:
            sys.exit("Cap dels identificadors demanats existeix. Res a fer.")

    files_per_carpeta = {}
    conf_per_carpeta = {}
    totes_les_files = []
    for c in carpetes:
        fitxers, conf = carrega_expedient(c)
        if fitxers:
            files_per_carpeta[c.name] = fitxers
            conf_per_carpeta[c.name] = conf
            totes_les_files.extend(fitxers)

    if not totes_les_files:
        sys.exit("Cap carpeta té files.jsonl amb contingut.")

    extres = columnes_extra_ordenades(totes_les_files)

    headers = [label for _, label in CORE] + [NOTES_LABEL] + extres + \
              ["Confiança", "Motiu confiança", "Font"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Extraccio"

    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"

    fila_excel = 2
    for nom_carpeta in files_per_carpeta:
        fitxers = files_per_carpeta[nom_carpeta]
        conf = conf_per_carpeta[nom_carpeta]
        for idx, f in enumerate(fitxers, start=1):
            col = 1
            for key, _ in CORE:
                v = f.get(key, "")
                ws.cell(row=fila_excel, column=col, value=v).font = BODY_FONT
                col += 1
            ws.cell(row=fila_excel, column=col, value=f.get(NOTES_KEY, "")).font = BODY_FONT
            col += 1
            for k in extres:
                ws.cell(row=fila_excel, column=col, value=f.get(k, "")).font = BODY_FONT
                col += 1

            info = conf.get(idx, {})
            c_val = info.get("confianca", "")
            c_cell = ws.cell(row=fila_excel, column=col, value=c_val)
            c_cell.font = Font(name="Arial", size=10, bold=True,
                                color=CONF_COLOR.get(c_val, "000000"))
            col += 1
            ws.cell(row=fila_excel, column=col, value=info.get("motiu", "")).font = BODY_FONT
            col += 1
            ws.cell(row=fila_excel, column=col, value=info.get("font", "")).font = BODY_FONT

            if fila_excel % 2 == 1:
                for cc in range(1, len(headers) + 1):
                    if ws.cell(row=fila_excel, column=cc).fill.fgColor.rgb in (None, "00000000"):
                        ws.cell(row=fila_excel, column=cc).fill = ALT_FILL

            fila_excel += 1

    amples = {
        "Descripció": 42, "Notes": 32, "Motiu confiança": 40, "Font": 30,
        "Proveïdor": 26, "Carpeta analitzada": 14, "Expedient": 14,
    }
    for col, text in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col)].width = amples.get(text, 16)

    wb.save(sortida)
    print(f"Escrit {sortida}  —  {fila_excel - 2} files de {len(files_per_carpeta)} expedient(s): "
          f"{', '.join(files_per_carpeta)}")


if __name__ == "__main__":
    main()

'''
python build_excel.py                                              # tots
python build_excel.py output extraccio_licitacions.xlsx all        # tots, explícit
python build_excel.py output sonnet_ok.xlsx 4161038,9252209        # només aquests dos
python build_excel.py output opus_repesca.xlsx "2222 3333 4444"    # també amb espais
'''